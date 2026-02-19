#!/usr/bin/env python3
import json
import sys
from pathlib import Path


def fallback_lines(text: str):
    return [p.strip() for p in text.split("\n") if p.strip()]


def parse_txt(path: Path):
    data = path.read_text(errors="ignore")
    lines = fallback_lines(data)
    return [{"page": 1, "text": "\n".join(lines)}]


def parse_pdf(path: Path):
    pages = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        for i, page in enumerate(reader.pages, start=1):
            txt = (page.extract_text() or "").strip()
            pages.append({"page": i, "text": txt})
    except Exception:
        pages = []
    if not pages:
        pages = [{"page": 1, "text": ""}]
    return pages


def parse_docx(path: Path):
    pages = []
    try:
        import docx

        d = docx.Document(str(path))
        text = "\n".join([p.text for p in d.paragraphs if p.text.strip()])
        pages = [{"page": 1, "text": text}]
    except Exception:
        pages = [{"page": 1, "text": ""}]
    return pages


def parse_xlsx(path: Path):
    pages = []
    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        chunks = []
        for ws in wb.worksheets:
            chunks.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(c) for c in row if c is not None and str(c).strip()]
                if vals:
                    chunks.append(" | ".join(vals))
        pages = [{"page": 1, "text": "\n".join(chunks)}]
    except Exception:
        pages = [{"page": 1, "text": ""}]
    return pages


def main():
    if len(sys.argv) != 3:
        print("Usage: parse_docs.py <source_file> <parsed_root>")
        sys.exit(1)

    src = Path(sys.argv[1])
    parsed_root = Path(sys.argv[2])
    pages_root = parsed_root / "pages"
    parsed_root.mkdir(parents=True, exist_ok=True)
    pages_root.mkdir(parents=True, exist_ok=True)

    ext = src.suffix.lower()
    if ext == ".pdf":
        pages = parse_pdf(src)
    elif ext == ".docx":
        pages = parse_docx(src)
    elif ext == ".xlsx":
        pages = parse_xlsx(src)
    else:
        pages = parse_txt(src)

    full_text = []
    for item in pages:
        p = item["page"]
        t = item["text"]
        full_text.append(f"\n\n--- PAGE {p} ---\n{t}")
        (pages_root / f"{src.name}_p{p}.txt").write_text(t, encoding="utf-8")

    (parsed_root / f"{src.name}.txt").write_text("".join(full_text), encoding="utf-8")

    print(json.dumps({"file": src.name, "page_count": len(pages)}))


if __name__ == "__main__":
    main()
